from langchain_core.tools import tool
from docx import Document
from reportlab.pdfgen import canvas
import os

@tool
def create_brand_strategy_document(strategy_content: dict, output_path: str = "output/strategy.docx") -> str:
    """
    Creates a Word document containing the brand strategy.
    """
    print(f"📄 [DOCUMENT TOOL] Generating strategy document at {output_path}...")
    
    doc = Document()
    doc.add_heading('Brand Strategy', 0)
    
    if 'positioning' in strategy_content:
        doc.add_heading('Positioning', level=1)
        doc.add_paragraph(strategy_content['positioning'])
    
    if 'personality' in strategy_content:
        doc.add_heading('Brand Personality', level=1)
        for trait in strategy_content['personality']:
            doc.add_paragraph(f"- {trait}", style='List Bullet')

    if 'messaging_framework' in strategy_content:
        doc.add_heading('Messaging Framework', level=1)
        for key, value in strategy_content['messaging_framework'].items():
            doc.add_paragraph(f"{key}: {value}")

    # Ensure output directory exists
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path

@tool
def create_brand_guidelines_pdf(design_outputs: dict, output_path: str = "output/guidelines.pdf") -> str:
    """
    Creates a simple PDF for brand guidelines.
    """
    print(f"📄 [DOCUMENT TOOL] Generating brand guidelines PDF at {output_path}...")
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    c = canvas.Canvas(output_path)
    c.drawString(100, 800, "Brand Guidelines")
    
    y = 750
    if 'color_palette' in design_outputs:
        c.drawString(100, y, "Color Palette:")
        y -= 20
        for color in design_outputs['color_palette']:
            c.drawString(120, y, f"- {color}")
            y -= 15
            
    if 'typography' in design_outputs:
        y -= 20
        c.drawString(100, y, "Typography:")
        y -= 20
        type_data = design_outputs['typography']
        # Handle if it's a string or dict
        if isinstance(type_data, dict):
            for k, v in type_data.items():
                c.drawString(120, y, f"{k}: {v}")
                y -= 15
        else:
             c.drawString(120, y, str(type_data))
    
    c.save()
    return output_path
